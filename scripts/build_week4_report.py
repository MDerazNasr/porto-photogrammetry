#!/usr/bin/env python3
"""Build Mohamed's Fall 2026 Week 4 photogrammetry report."""

from pathlib import Path

from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, ListFlowable, ListItem, PageBreak, Paragraph
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
PDF_PATH = ROOT / "output/pdf/Week 4 Report - Mohamed.pdf"
DOCX_PATH = ROOT / "output/docx/Week 4 Report - Mohamed.docx"
METRICS_IMAGE = ROOT / "outputs/week4-color-calibration-metrics.png"
COMPARISON_IMAGE = ROOT / "outputs/color-calibration-batch-real-images-2026-09-22.jpg"

IDENTITY = ["Mohamed Deraz Nasr", "HAAG Photogrammetry Project", "Fall 2026", "September 25, 2026"]

PROGRESS = [
    "Opened upstream PR 18, Add per-camera ColorChecker calibration command, from the forked mohamed/color-calibration branch and packaged a reproducible synthetic demonstration.",
    "Added confidence-gated automatic chart detection while retaining reviewed four-corner input as the safe fallback for contact, occlusion, or low-confidence cases.",
    "Added colour-target provenance so the command supports relative camera alignment now and a defensible absolute target once the physical chart edition is confirmed.",
    "Implemented specimen-mask-only clipping diagnostics and a conservative 0.5-percentage-point warning threshold; both observed archive mask-name conventions are supported.",
    "Repeated the method on UF_birds_ivory2: 138 full-resolution images across three usable cameras, with the partly occluded fourth chart explicitly excluded.",
    "Ran an independent 36-view original-versus-calibrated reconstruction and CPU sparse-Delaunay mesh comparison after full dense MVS was blocked by unavailable CUDA/PACE access.",
    "Expanded focused coverage to 14 passing colour tests and updated the PR, CLI documentation, evidence reports, weekly database, and handoff notes.",
]

RESULT_ROWS = [
    ("UF_Herp_3998 camera 2 Delta E76", "9.218", "3.131", "-66.0%"),
    ("UF_Herp_3998 camera 3 Delta E76", "3.698", "1.432", "-61.3%"),
    ("UF_birds_ivory2 camera 2 Delta E76", "7.998", "2.866", "-64.2%"),
    ("UF_birds_ivory2 camera 3 Delta E76", "7.380", "1.614", "-78.1%"),
    ("Second-specimen registrations", "33 / 36", "33 / 36", "Unchanged"),
    ("Sparse-Delaunay mesh vertices", "3,506", "3,476", "-0.86%"),
    ("Sparse-Delaunay mesh faces", "6,962", "6,905", "-0.82%"),
    ("Mean reprojection error", "0.8510 px", "0.8530 px", "+0.0020 px"),
]

BLOCKERS = [
    "Absolute calibration remains blocked on the physical ColorChecker manufacturer/model and patch-value edition. The software path is ready, but the current validated claim remains relative camera alignment.",
    "Camera 2 on UF_birds_ivory2 increased foreground clipping from 0.050% to 0.612%, crossing the warning threshold. More specimens are needed before this threshold becomes an acceptance gate.",
    "The available local COLMAP build requires CUDA for PatchMatch and PACE was unreachable, so the mesh result is sparse-Delaunay topology evidence, not dense-MVS or anatomical accuracy.",
]

DOCUMENTATION = [
    ("Pilot and decision record", "documentation/color-calibration-pilot-2026-09-21.md"),
    ("Implementation", "team-repo/src/augenblick/preparation/color.py and augenblick color"),
    ("Second-specimen evidence", "outputs/color-calibration-report-uf-birds-ivory2-2026-09-25.json"),
    ("Mesh regression", "outputs/color-calibration-mesh-regression-uf-birds-ivory2-2026-09-25.json"),
    ("Mesh comparison utility", "scripts/compare_colmap_meshes.py"),
    ("Pull request", "Human-Augment-Analytics/porto-photogrammetry PR 18"),
    ("Week 5 handoff", "documentation/week-5-todo-2026-09-28.md"),
]

NEXT_WEEK = [
    "Obtain the physical chart model and patch-value edition, then rerun the supported absolute-target mode and finalize that part of PR 18.",
    "Coordinate coded-marker scaling with Ihor's accepted windowed protocol; contribute integration, an independent specimen, or colour/scale interaction testing rather than duplicating his work.",
    "Run a dense reconstruction or reference-mesh geometry comparison when suitable CUDA/PACE resources and ground truth are available.",
    "Compare the 3 by 3 model with root-polynomial correction on held-out patches only; require a clear gain without clipping, casts, or exposure instability.",
    "Request pipeline-owner and Om review, post the PR/demo and evidence in Slack, and turn feedback into named follow-up issues.",
    "Review Ihor Week 4/5, Om Week 4, and Syed Week 4 for overlaps, reusable artifacts, and coordination decisions.",
]

VERIFICATION = [
    ("PR and demo", "Opened PR 18 and documented a sanitized, reproducible command path."),
    ("Detection safety", "Automatic detection declines uncertain/contact cases; reviewed corners remain available."),
    ("Mask compatibility", "Supports both <stem>.mask.png and <image-name>.mask.png; 14 focused tests pass."),
    ("Independent colour run", "Processed 138 usable images; camera 4 was skipped because its chart was occluded."),
    ("Clipping gate", "Foreground-only metric raised one explicit warning instead of hiding background-dominated clipping."),
    ("Geometry regression", "Both runs registered 33/36; mesh size changed by under 1% and aligned mean vertex distance was 0.169% of the robust bounding-box diagonal."),
]

ABSTRACTS = [
    (
        "Preserving colour fidelity in photogrammetry",
        "Barbero-Alvarez et al. study how acquisition, RAW development, colour space, chart calibration, and texture generation affect cultural-heritage photogrammetry. Their central recommendation is to control colour transformations and evaluate known chart patches after processing. This supports treating colour fidelity as a measurable pipeline property rather than relying on visual appearance alone.",
        "Barbero-Alvarez, M. A., Brenner, S., Sablatnig, R., and Menendez, J. M. (2023). Heritage, 6(8), 5700-5718. https://doi.org/10.3390/heritage6080300",
    ),
    (
        "Color correction using root-polynomial regression",
        "Finlayson et al. compare standard 3 by 3 linear colour correction with polynomial alternatives. Linear correction is simple and exposure-independent, while ordinary polynomial models may introduce exposure-dependent hue shifts; root-polynomial terms retain exposure scaling and can improve accuracy. The present pilot deliberately begins with the lower-complexity linear model and leaves higher-order fitting for held-out evaluation.",
        "Finlayson, G. D., Mackiewicz, M., and Hurlbert, A. (2015). IEEE Transactions on Image Processing, 24(5), 1460-1470. https://doi.org/10.1109/TIP.2015.2405336",
    ),
    (
        "Imageomics colour-calibration software",
        "The Imageomics package uses segmented colour cards and per-channel histogram matching to normalize target images to a reference image. Our real-data test confirmed that histogram convergence does not guarantee patch accuracy or natural full-image colour, motivating the corresponding-patch transform implemented this week.",
        "Imageomics Institute. color-calibration software repository. https://github.com/Imageomics/color-calibration (accessed September 23, 2026).",
    ),
]

NAVY, LIGHT_GREY, MID_GREY = "1F4E79", "F2F2F2", "666666"


def font(size, bold=False):
    names = ["Arial Bold.ttf", "Arial.ttf"] if bold else ["Arial.ttf"]
    for name in names:
        for base in ("/System/Library/Fonts/Supplemental", "/Library/Fonts"):
            path = Path(base) / name
            if path.exists():
                return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def build_metrics_image():
    canvas = PILImage.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(canvas)
    navy, blue, orange, grey = (31, 78, 121), (91, 155, 213), (237, 125, 49), (95, 95, 95)
    draw.text((60, 38), "Colour calibration validation", font=font(50, True), fill=navy)
    draw.line((60, 110, 1740, 110), fill=navy, width=4)

    draw.rounded_rectangle((60, 145, 865, 815), radius=22, fill=(235, 241, 247))
    draw.text((95, 175), "Mean patch colour error (Delta E76)", font=font(38, True), fill=navy)
    draw.text((95, 225), "Lower is better", font=font(28), fill=grey)
    for idx, (label, before, after) in enumerate((("Herp C2", 9.218, 3.131), ("Herp C3", 3.698, 1.432), ("Bird C2", 7.998, 2.866), ("Bird C3", 7.380, 1.614))):
        y, x0, scale = 290 + idx * 118, 285, 45
        draw.text((95, y + 18), label, font=font(29), fill=(30, 30, 30))
        draw.rounded_rectangle((x0, y, x0 + before * scale, y + 48), radius=10, fill=orange)
        draw.rounded_rectangle((x0, y + 54, x0 + after * scale, y + 96), radius=10, fill=blue)
        draw.text((x0 + before * scale + 15, y + 4), f"{before:.3f}", font=font(28), fill=(30, 30, 30))
        draw.text((x0 + after * scale + 15, y + 56), f"{after:.3f}", font=font(28), fill=(30, 30, 30))
    draw.rectangle((95, 770, 125, 800), fill=orange)
    draw.text((140, 768), "Before", font=font(28), fill=grey)
    draw.rectangle((300, 770, 330, 800), fill=blue)
    draw.text((345, 768), "After", font=font(28), fill=grey)

    draw.rounded_rectangle((905, 145, 1740, 815), radius=22, fill=(247, 247, 247))
    draw.text((940, 175), "Second-specimen regression", font=font(38, True), fill=navy)
    draw.text((940, 225), "Calibrated relative to original", font=font(28), fill=grey)
    rows = (("Registered views", "33/36 -> 33/36"), ("Sparse points", "5,008 -> 4,999"), ("Reprojection", "0.8510 -> 0.8530 px"), ("Mesh vertices", "3,506 -> 3,476"), ("Mesh faces", "6,962 -> 6,905"))
    for idx, (label, value) in enumerate(rows):
        y = 300 + idx * 78
        draw.text((940, y), label, font=font(27), fill=(30, 30, 30))
        draw.text((1250, y), value, font=font(27, True), fill=navy)
    draw.text((940, 710), "Aligned symmetric mean vertex distance", font=font(26), fill=grey)
    draw.text((940, 752), "0.169% of robust bbox diagonal", font=font(31, True), fill=navy)
    METRICS_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(METRICS_IMAGE, optimize=True)


def set_run(run, size=10.2, bold=False, italic=False, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_format(p, before=0, after=0, line=1.08):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if node.getparent() is None:
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margin(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tc_mar.getparent() is None:
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
        if node.getparent() is None:
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def identity_docx(doc):
    for line in IDENTITY:
        p = doc.add_paragraph(); paragraph_format(p); set_run(p.add_run(line), 9.4)


def heading_docx(doc, text, size=18, before=13, after=5):
    p = doc.add_paragraph(); paragraph_format(p, before, after); set_run(p.add_run(text), size, color=NAVY)


def bullet_docx(doc, text):
    p = doc.add_paragraph(style="List Bullet"); paragraph_format(p, after=1); set_run(p.add_run(text), 10)


def table_docx(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader"); tbl_header.set(qn("w:val"), "true"); tr_pr.append(tbl_header)
    for i, (cell, label, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(cell); shade(cell, NAVY)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_run(cell.paragraphs[0].add_run(label), 8.8, True, color="FFFFFF")
    for row_no, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = width; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(cell)
            if row_no % 2: shade(cell, LIGHT_GREY)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run(cell.paragraphs[0].add_run(value), 8.5)
    return table


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def caption_docx(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; paragraph_format(p, 2, 7, 1.0)
    set_run(p.add_run(text), 8.5, italic=True, color=MID_GREY)


def build_docx():
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.78); sec.top_margin, sec.bottom_margin = Inches(0.5), Inches(0.55)
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(10.2); normal.paragraph_format.space_after = Pt(0)
    bullet = doc.styles["List Bullet"]; bullet.paragraph_format.left_indent = Inches(0.42); bullet.paragraph_format.first_line_indent = Inches(-0.22)

    identity_docx(doc)
    p = doc.add_paragraph(); paragraph_format(p, 22, 2); set_run(p.add_run("Week 4 Report"), 26, color=NAVY)
    p = doc.add_paragraph(); paragraph_format(p, after=10); set_run(p.add_run("Colour calibration implementation and validation"), 12, color=MID_GREY)
    heading_docx(doc, "Progress this week", 16, 2)
    for x in PROGRESS: bullet_docx(doc, x)
    heading_docx(doc, "Key results", 16, 9)
    table_docx(doc, ("Metric", "Original", "Calibrated", "Change"), RESULT_ROWS,
               (Inches(3.0), Inches(1.2), Inches(1.2), Inches(1.15)))
    heading_docx(doc, "Interpretation", 16, 10)
    p = doc.add_paragraph(); paragraph_format(p, after=4); set_run(p.add_run("Colour error fell on both specimens. The unchanged registration counts, sub-1% mesh-size changes, and small aligned surface difference support the narrower claim that calibration preserved coarse reconstruction behaviour. They do not prove dense or anatomical accuracy."), 10)
    heading_docx(doc, "Blockers", 16, 9)
    for x in BLOCKERS: bullet_docx(doc, x)

    page_break(doc); identity_docx(doc); heading_docx(doc, "2. Literature review", 20, 18)
    for title, summary, citation in ABSTRACTS:
        p = doc.add_paragraph(); paragraph_format(p, 8, 2); set_run(p.add_run(title), 11, True, color=NAVY)
        p = doc.add_paragraph(); paragraph_format(p, after=3); set_run(p.add_run(summary), 10)
        p = doc.add_paragraph(); paragraph_format(p, after=5, line=1.05); set_run(p.add_run("Citation: "), 8.7, True, color=MID_GREY)
        set_run(p.add_run(citation), 8.7, color=MID_GREY)

    page_break(doc); identity_docx(doc); heading_docx(doc, "3. What I did and how I verified it", 20, 18)
    for title, detail in VERIFICATION:
        p = doc.add_paragraph(); paragraph_format(p, 7, 1); set_run(p.add_run(title), 10.5, True, color=NAVY)
        p = doc.add_paragraph(); paragraph_format(p, after=3); set_run(p.add_run(detail), 10)

    page_break(doc); identity_docx(doc); heading_docx(doc, "4. Results visualization", 20, 18, 7)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(COMPARISON_IMAGE), width=Inches(4.65))
    caption_docx(doc, "Figure 1. Representative original and calibrated images from the three-camera batch.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(METRICS_IMAGE), width=Inches(6.7))
    caption_docx(doc, "Figure 2. Patch-error reduction across both specimens and the second-specimen sparse-Delaunay regression.")

    page_break(doc); identity_docx(doc); heading_docx(doc, "5. Conclusions and decisions", 20, 18)
    for x in ("Keep the corresponding-patch 3 by 3 linear method as the validated default.", "Use automatic detection only when confidence gates pass; keep reviewed corners as the fallback.", "Treat the camera-2 foreground clipping increase as a warning that requires review, not a hidden failure or a universal rejection threshold.", "Describe geometry as preserved at sparse/coarse topology level; do not claim dense or anatomical validation.", "Keep absolute colour calibration and metric scale as separate, provenance-controlled stages."):
        bullet_docx(doc, x)
    heading_docx(doc, "Evidence and artifacts", 16, 10)
    table_docx(doc, ("Artifact", "Location"), DOCUMENTATION, (Inches(2.05), Inches(4.5)))

    page_break(doc); identity_docx(doc); heading_docx(doc, "Next week", 20, 18)
    for x in NEXT_WEEK: bullet_docx(doc, x)
    heading_docx(doc, "Team coordination", 16, 12)
    p = doc.add_paragraph(); paragraph_format(p, after=4); set_run(p.add_run("The Slack database now includes Ihor Week 4 and 5, Om Week 4, and Syed Week 4. Ihor's coded-marker work is already beyond a starting prototype, so the next step is coordination and complementary validation. Meeting 4 also requires matched hardware for wall-clock comparisons and communication before resolving resource conflicts."), 10)
    alt_texts = [
        "Three rows of original and calibrated museum specimen photographs from cameras 1, 2, and 3.",
        "Chart showing Delta E76 reductions on two specimens and stable second-specimen sparse reconstruction and mesh metrics.",
    ]
    for shape, alt in zip(doc.inline_shapes, alt_texts):
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt); doc_pr.set("title", alt)
    doc.core_properties.title, doc.core_properties.author = "Week 4 Report", "Mohamed Deraz Nasr"
    doc.core_properties.subject = "HAAG Photogrammetry Project"
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True); doc.save(DOCX_PATH)


def styles_pdf():
    base = getSampleStyleSheet()["Normal"]
    body = ParagraphStyle("body", parent=base, fontName="Helvetica", fontSize=9.5, leading=11.8, alignment=TA_LEFT, textColor=colors.HexColor("#222222"))
    identity = ParagraphStyle("identity", parent=body, fontSize=9.3, leading=11, keepWithNext=1)
    title = ParagraphStyle("title", parent=body, fontSize=26, leading=30, spaceBefore=21, spaceAfter=2, textColor=colors.HexColor("#1F4E79"))
    subtitle = ParagraphStyle("subtitle", parent=body, fontSize=12, leading=15, spaceAfter=9, textColor=colors.HexColor("#666666"))
    heading = ParagraphStyle("heading", parent=body, fontSize=18, leading=22, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#1F4E79"))
    small_heading = ParagraphStyle("small_heading", parent=heading, fontSize=15, leading=18, spaceBefore=8, spaceAfter=4)
    caption = ParagraphStyle("caption", parent=body, fontSize=8.2, leading=10, alignment=TA_CENTER, textColor=colors.HexColor("#666666"), spaceBefore=2, spaceAfter=6)
    paper = ParagraphStyle("paper", parent=body, fontSize=11, leading=13, textColor=colors.HexColor("#1F4E79"), spaceBefore=8, spaceAfter=2)
    citation = ParagraphStyle("citation", parent=body, fontSize=8.4, leading=10.4, textColor=colors.HexColor("#555555"), leftIndent=10, spaceAfter=5)
    return body, identity, title, subtitle, heading, small_heading, caption, paper, citation


def identity_pdf(story, style):
    story.extend(Paragraph(x, style) for x in IDENTITY)


def bullets_pdf(items, body, size=9.4):
    item = ParagraphStyle(f"item_{size}", parent=body, fontSize=size, leading=size * 1.24, spaceAfter=1)
    return ListFlowable([ListItem(Paragraph(x, item), leftIndent=17) for x in items], bulletType="bullet",
                        bulletChar="bullet", leftIndent=17, bulletFontName="Helvetica", bulletFontSize=8.5)


def table_pdf(headers, rows, widths):
    table = Table([list(headers)] + [list(r) for r in rows], colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.2), ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7B7B7")), ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_pdf():
    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    body, ident, title, subtitle, heading, small_heading, caption, paper, citation = styles_pdf()
    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=letter, leftMargin=.78*inch, rightMargin=.78*inch,
                            topMargin=.48*inch, bottomMargin=.52*inch, title="Week 4 Report", author="Mohamed Deraz Nasr")
    story = []; identity_pdf(story, ident)
    story += [Paragraph("Week 4 Report", title), Paragraph("Colour calibration implementation and validation", subtitle),
              Paragraph("Progress this week", small_heading), bullets_pdf(PROGRESS, body, 9.25),
              Paragraph("Key results", small_heading),
              table_pdf(("Metric", "Original", "Calibrated", "Change"), RESULT_ROWS, [2.85*inch, 1.18*inch, 1.18*inch, 1.15*inch]),
              Paragraph("Interpretation", small_heading),
              Paragraph("Colour error fell on both specimens. The unchanged registration counts, sub-1% mesh-size changes, and small aligned surface difference support the narrower claim that calibration preserved coarse reconstruction behaviour. They do not prove dense or anatomical accuracy.", body),
              Paragraph("Blockers", small_heading), bullets_pdf(BLOCKERS, body, 9.1)]

    story.append(PageBreak()); identity_pdf(story, ident); story.append(Paragraph("2. Literature review", heading))
    for paper_title, summary, source in ABSTRACTS:
        story += [Paragraph(paper_title, paper), Paragraph(summary, body), Paragraph(f"<b>Citation:</b> {source}", citation)]

    story.append(PageBreak()); identity_pdf(story, ident); story.append(Paragraph("3. What I did and how I verified it", heading))
    for item_title, detail in VERIFICATION:
        story += [Paragraph(item_title, paper), Paragraph(detail, body)]

    story.append(PageBreak()); identity_pdf(story, ident); story.append(Paragraph("4. Results visualization", heading))
    story += [Image(str(COMPARISON_IMAGE), 4.62*inch, 4.62*inch),
              Paragraph("Figure 1. Representative original and calibrated images from the three-camera batch.", caption),
              Image(str(METRICS_IMAGE), 6.72*inch, 3.36*inch),
              Paragraph("Figure 2. Patch-error reduction across both specimens and the second-specimen sparse-Delaunay regression.", caption)]

    story.append(PageBreak()); identity_pdf(story, ident); story.append(Paragraph("5. Conclusions and decisions", heading))
    conclusions = ["Keep the corresponding-patch 3 by 3 linear method as the validated default.", "Use automatic detection only when confidence gates pass; keep reviewed corners as the fallback.", "Treat the camera-2 foreground clipping increase as a warning requiring review, not a universal rejection threshold.", "Describe geometry as preserved at sparse/coarse topology level; do not claim dense or anatomical validation.", "Keep absolute colour calibration and metric scale as separate, provenance-controlled stages."]
    story += [bullets_pdf(conclusions, body, 9.4), Paragraph("Evidence and artifacts", small_heading),
              table_pdf(("Artifact", "Location"), DOCUMENTATION, [2.0*inch, 4.35*inch])]

    story.append(PageBreak()); identity_pdf(story, ident); story.append(Paragraph("Next week", heading))
    story += [bullets_pdf(NEXT_WEEK, body, 9.4), Paragraph("Team coordination", small_heading),
              Paragraph("The Slack database now includes Ihor Week 4 and 5, Om Week 4, and Syed Week 4. Ihor's coded-marker work is already beyond a starting prototype, so the next step is coordination and complementary validation. Meeting 4 also requires matched hardware for wall-clock comparisons and communication before resolving resource conflicts.", body)]
    doc.build(story)


def main():
    if not COMPARISON_IMAGE.exists(): raise FileNotFoundError(COMPARISON_IMAGE)
    build_metrics_image(); build_docx(); build_pdf()
    print(DOCX_PATH); print(PDF_PATH); print(METRICS_IMAGE)


if __name__ == "__main__": main()
