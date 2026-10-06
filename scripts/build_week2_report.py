#!/usr/bin/env python3
"""Build the Week 2 photogrammetry report as matching DOCX and PDF files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
DOCX_PATH = OUT / "week-2-photogrammetry-report-2026-09-11.docx"
PDF_PATH = OUT / "week-2-photogrammetry-report-2026-09-11.pdf"

IDENTITY = [
    "Mohamed Deraz Nasr",
    "HAAG Photogrammetry Project",
    "Fall 2026",
    "September 11",
]

PROGRESS = [
    ("Reviewed the current Augenblick architecture and traced image preparation, pose estimation, reconstruction, mesh export, and evaluation", []),
    ("Confirmed that all reconstruction methods consume a common COLMAP scene containing images, optional masks, and a sparse camera model", []),
    ("Reviewed the Linux, CUDA, PACE, and HiPerGator environment requirements", []),
    ("Audited the current color pipeline: images are not ColorChecker-calibrated, and existing exposure compensation and appearance metrics do not establish true color accuracy", []),
    ("Confirmed that coded-marker detection and conversion to millimetres are not currently implemented", []),
    ("Reviewed color-calibration literature and identified ColorChecker-based correction with CIEDE2000 evaluation as the appropriate direction", []),
    ("Reviewed the Imageomics package: it uses COCO card masks and RGB histogram matching to improve internal consistency, but does not use official ColorChecker values", []),
    ("Proposed a small first deliverable: calibrate one capture session, preserve the original images, and report color error before and after correction", []),
]

NEXT_WEEK = [
    "Arrange a working session with Om and divide the investigation clearly so that we do not duplicate effort.",
    "Compare the existing Imageomics color-calibration tool's inputs and assumptions with the museum images and the current Augenblick workflow.",
    "Confirm the ColorChecker model and access to representative photographs, chart regions, segmentation masks, and the required execution environment.",
    "Run one small representative calibration test if the required data and environment are available.",
    "Document the adaptations, missing inputs, blockers, and likely integration points, and contact Arthur midweek if a decision is needed.",
    "Bring a clear recommendation to the next meeting: adapt the existing tool, or explain why it is unsuitable and propose the next alternative.",
]

ABSTRACTS = [
    (
        "Preserving Colour Fidelity in Photogrammetry - An Empirically Grounded Study and Workflow for Cultural Heritage Preservation",
        "This study examines how image acquisition, RAW development, color-space selection, ColorChecker calibration, and photogrammetric processing affect the color fidelity of cultural-heritage 3D models. It recommends controlling image transformations, calibrating photographs with known chart values, and evaluating color differences with CIEDE2000. The paper supports treating color accuracy as a separate measurement problem rather than relying only on visual similarity or reconstruction metrics.",
    ),
    (
        "Imageomics Color Calibration Pipeline",
        "The Imageomics package calibrates a dataset by choosing one photograph as a reference, extracting the color card from each annotated image, and matching the target card's red, green, and blue histograms to the reference card. The resulting lookup tables are applied to each full photograph and saved as PNG files. This can reduce variation between images, but it provides relative normalization rather than absolute ColorChecker calibration because it does not use official patch values or report CIEDE2000 error.",
    ),
]


def set_font(run, size=11, bold=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold


def configure_paragraph(paragraph, before=0, after=0, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_docx_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    configure_paragraph(p, after=0, line=1.15)
    set_font(p.add_run(text))
    return p


def configure_numbering_styles(doc):
    for style_name, left, first in (("List Bullet", 0.50, -0.25), ("List Bullet 2", 0.82, -0.20)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(left)
        style.paragraph_format.first_line_indent = Inches(first)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.15


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    configure_numbering_styles(doc)

    for line in IDENTITY:
        p = doc.add_paragraph()
        configure_paragraph(p)
        set_font(p.add_run(line))

    p = doc.add_paragraph()
    configure_paragraph(p, before=35, after=15)
    set_font(p.add_run("Week 2 Report"), size=26)

    p = doc.add_paragraph()
    configure_paragraph(p, after=5)
    set_font(p.add_run("Progress this week:"), size=16)
    for text, children in PROGRESS:
        add_docx_bullet(doc, text)
        for child in children:
            add_docx_bullet(doc, child, 1)

    p = doc.add_paragraph()
    configure_paragraph(p, before=16, after=5)
    set_font(p.add_run("Next Week"), size=20)
    for item in NEXT_WEEK:
        add_docx_bullet(doc, item)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    for line in IDENTITY:
        p = doc.add_paragraph()
        configure_paragraph(p)
        set_font(p.add_run(line))

    p = doc.add_paragraph()
    configure_paragraph(p, before=30, after=8)
    set_font(p.add_run("Abstracts"), size=20)
    for title, summary in ABSTRACTS:
        p = doc.add_paragraph()
        configure_paragraph(p, before=8, after=2)
        set_font(p.add_run(title), size=11)
        add_docx_bullet(doc, summary)

    doc.core_properties.title = "Week 2 Report"
    doc.core_properties.author = "Mohamed Deraz Nasr"
    doc.core_properties.subject = "HAAG Photogrammetry Project"
    doc.save(DOCX_PATH)


def build_pdf():
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=11,
                          leading=14.2, alignment=TA_LEFT, spaceAfter=0)
    identity = ParagraphStyle("identity", parent=body, leading=14.2)
    title = ParagraphStyle("title", parent=body, fontSize=26, leading=31, spaceBefore=34, spaceAfter=17)
    section = ParagraphStyle("section", parent=body, fontSize=20, leading=24, spaceBefore=16, spaceAfter=8)
    progress_heading = ParagraphStyle("progress", parent=body, fontSize=16, leading=20, spaceAfter=5)
    item_style = ParagraphStyle("item", parent=body, leftIndent=0, firstLineIndent=0)
    subitem_style = ParagraphStyle("subitem", parent=body, leftIndent=0, firstLineIndent=0)
    paper_title = ParagraphStyle("paper_title", parent=body, spaceBefore=8, spaceAfter=2)

    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=letter, leftMargin=inch, rightMargin=inch,
                            topMargin=0.52 * inch, bottomMargin=0.7 * inch,
                            title="Week 2 Report", author="Mohamed Deraz Nasr")
    story = [Paragraph(line, identity) for line in IDENTITY]
    story += [Paragraph("Week 2 Report", title), Paragraph("Progress this week:", progress_heading)]

    progress_items = []
    for text, children in PROGRESS:
        content = [Paragraph(text, item_style)]
        if children:
            content.append(ListFlowable(
                [ListItem(Paragraph(c, subitem_style)) for c in children],
                bulletType="bullet", bulletChar="○", leftIndent=24, bulletFontName="Helvetica",
                bulletFontSize=9, spaceBefore=2, spaceAfter=1,
            ))
        progress_items.append(ListItem(content, leftIndent=18))
    story.append(ListFlowable(progress_items, bulletType="bullet", bulletChar="•", leftIndent=18,
                              bulletFontName="Helvetica", bulletFontSize=11, spaceAfter=4))

    story.append(Paragraph("Next Week", section))
    story.append(ListFlowable(
        [ListItem(Paragraph(x, item_style), leftIndent=18) for x in NEXT_WEEK],
        bulletType="bullet", bulletChar="•", leftIndent=18, bulletFontName="Helvetica", bulletFontSize=11,
    ))

    story.append(PageBreak())
    story.extend(Paragraph(line, identity) for line in IDENTITY)
    story.append(Paragraph("Abstracts", section))
    for paper, summary in ABSTRACTS:
        story.append(Paragraph(paper, paper_title))
        story.append(ListFlowable([ListItem(Paragraph(summary, item_style), leftIndent=18)],
                                  bulletType="bullet", bulletChar="•", leftIndent=18,
                                  bulletFontName="Helvetica", bulletFontSize=11))
    doc.build(story)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
